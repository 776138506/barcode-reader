"""模板渲染管线 v2（两段式）+ TXT/CSV/XLSX/JSON 导出 + 导出过滤。

行模板占位符（单码渲染）：
  {index}    序号（全局递增，从 1 开始，分组内不重启）
  {filename} 来源文件名（去重记录为分号连接的来源列表）
  {type}     码制
  {content}  码内容
  {count}    出现次数（去重记录 >1；普通记录恒为 1）
  {date}     导出日期（YYYY-MM-DD）
  {time}     导出时间（HH:MM:SS）

两段式（D16）：行模板 → 组内按连接符拼接 → 外模板（含 {items} 占位符）包装。
外模板为 "{items}" 时退化为逐行行为（向后兼容旧配置）。
分组维度：none（不分组）/ image（按来源图聚合）/ global（全局聚合一条）。

未知占位符原样保留并记录警告。
"""
from __future__ import annotations

import csv
import io
import json
import re
import warnings
from dataclasses import dataclass, field
from datetime import datetime

DEFAULT_TEMPLATE = "{index},{filename},{type},{content}"
DEFAULT_JOINER = "\n"
DEFAULT_OUTER = "{items}"
DEFAULT_HEADER = "序号,文件名,码制,内容"

GROUP_BY = ("none", "image", "global")

KNOWN_PLACEHOLDERS = {"index", "filename", "type", "content", "count", "date", "time"}
KNOWN_OUTER_PLACEHOLDERS = {"items"}

_PLACEHOLDER_RE = re.compile(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}")


@dataclass
class ExportRecord:
    filename: str
    type: str
    content: str
    count: int = 1
    suspect: bool = False


@dataclass
class ExportFilter:
    """导出过滤（四条件可组合；None/空 = 不限制）。"""
    types: list[str] | None = None      # 选中的码制（显示名，None/空=全部）
    min_len: int | None = None          # 内容长度下限
    max_len: int | None = None          # 内容长度上限
    prefix: str = ""                    # 内容前缀
    regex: str = ""                     # 内容正则（re.search）


class _SafeDict(dict):
    """未知占位符原样保留（如 {foo} -> "{foo}"），不抛 KeyError。"""

    def __missing__(self, key):
        return "{" + key + "}"


def warn_unknown_placeholders(template: str, known: set | None = None) -> list[str]:
    """返回模板中的未知占位符列表，并发出警告。"""
    unknown = sorted(set(_PLACEHOLDER_RE.findall(template))
                     - (known or KNOWN_PLACEHOLDERS))
    if unknown:
        warnings.warn(
            "模板包含未知占位符（将原样保留）: " + ", ".join("{%s}" % u for u in unknown),
            UserWarning,
            stacklevel=2,
        )
    return unknown


def _build_context(record: ExportRecord, index: int, now: datetime) -> _SafeDict:
    return _SafeDict(
        index=index,
        filename=record.filename,
        type=record.type,
        content=record.content,
        count=record.count,
        date=now.strftime("%Y-%m-%d"),
        time=now.strftime("%H:%M:%S"),
    )


def render_template(template: str, record: ExportRecord, index: int,
                    now: datetime | None = None) -> str:
    """按行模板渲染单条记录。"""
    now = now or datetime.now()
    return template.format_map(_build_context(record, index, now))


# ---------------------------------------------------------------- 过滤

def apply_filter(records: list[ExportRecord], filt: ExportFilter | None) -> list[ExportRecord]:
    """应用过滤条件（可组合）。正则非法抛 ValueError（调用方负责提示）。"""
    if filt is None:
        return records
    pattern = None
    if filt.regex:
        try:
            pattern = re.compile(filt.regex)
        except re.error as exc:
            raise ValueError(f"无效正则 {filt.regex!r}: {exc}") from exc
    out = []
    for r in records:
        if filt.types and r.type not in filt.types:
            continue
        if filt.min_len is not None and len(r.content) < filt.min_len:
            continue
        if filt.max_len is not None and len(r.content) > filt.max_len:
            continue
        if filt.prefix and not r.content.startswith(filt.prefix):
            continue
        if pattern is not None and not pattern.search(r.content):
            continue
        out.append(r)
    return out


# ---------------------------------------------------------------- 分组与两段式渲染

def group_records(records: list[ExportRecord],
                  group_by: str = "none") -> list[tuple[str | None, list[ExportRecord]]]:
    """分组：none 每码一组（键为 None）；image 按来源文件名聚合；global 全部一组。"""
    if group_by == "image":
        groups: dict[str, list[ExportRecord]] = {}
        for r in records:
            groups.setdefault(r.filename, []).append(r)
        return list(groups.items())
    if group_by == "global":
        return [(None, records)] if records else []
    return [(None, [r]) for r in records]


def render_two_stage(records: list[ExportRecord],
                     template: str = DEFAULT_TEMPLATE,
                     joiner: str = DEFAULT_JOINER,
                     outer: str = DEFAULT_OUTER,
                     group_by: str = "none",
                     now: datetime | None = None) -> str:
    """两段式渲染：行模板逐码渲染 → 组内连接符拼接 → 外模板包装，组间换行。

    外模板只做字面 {items} 替换（不走 format 解析），因此其余字符
    （引号、花括号等）全部按字面输出，如 outer="{'{items}'}" → "{'abc','def'}"。
    outer 为 "{items}" 且 group_by 为 none 时等价于旧的逐行行为。
    """
    now = now or datetime.now()
    warn_unknown_placeholders(template)
    if "{items}" not in outer:
        warnings.warn("外模板缺少 {items} 占位符，渲染结果将不含码内容",
                      UserWarning, stacklevel=2)
    out_lines = []
    index = 0
    for _key, group in group_records(records, group_by):
        items = joiner.join(
            render_template(template, r, (index := index + 1), now) for r in group)
        out_lines.append(outer.replace("{items}", items))
    return "\n".join(out_lines)


# ---------------------------------------------------------------- 各格式写出

def export_txt(records: list[ExportRecord], path: str,
               template: str = DEFAULT_TEMPLATE,
               joiner: str = DEFAULT_JOINER,
               outer: str = DEFAULT_OUTER,
               group_by: str = "none") -> None:
    """TXT：两段式渲染整篇写出。默认参数等价旧版逐行行为。"""
    text = render_two_stage(records, template, joiner, outer, group_by)
    # newline=None（默认）：Windows 自动写出 CRLF、macOS/Linux 写出 LF，
    # 符合各平台文本编辑器的换行习惯
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)
        if text:
            f.write("\n")


def _split_columns(template: str) -> tuple[list[str], str]:
    """按分隔符（Tab 优先，其次逗号）把模板拆成列，返回 (列模板列表, 分隔符)。"""
    if "\t" in template:
        return template.split("\t"), "\t"
    return template.split(","), ","


def export_csv(records: list[ExportRecord], path: str,
               template: str = DEFAULT_TEMPLATE,
               header: str | None = None) -> None:
    """CSV：模板按分隔符拆列渲染，header 非空时写入首行表头。"""
    warn_unknown_placeholders(template)
    columns, delim = _split_columns(template)
    now = datetime.now()
    # csv 模块要求 newline=""，由 csv.writer 自己控制行尾（默认 \r\n），
    # utf-8-sig 带 BOM 方便 Windows 版 Excel 正确识别编码
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=delim)
        if header:
            writer.writerow(header.split(","))
        for i, r in enumerate(records):
            ctx = _build_context(r, i + 1, now)
            writer.writerow([col.format_map(ctx) for col in columns])


def export_xlsx(records: list[ExportRecord], path: str,
                template: str = DEFAULT_TEMPLATE,
                header: str | None = None) -> None:
    """XLSX：与 CSV 同语义的列渲染（openpyxl，纯 Python 三平台通用）。"""
    from openpyxl import Workbook
    warn_unknown_placeholders(template)
    columns, _delim = _split_columns(template)
    now = datetime.now()
    wb = Workbook()
    ws = wb.active
    ws.title = "barcodes"
    if header:
        ws.append(header.split(","))
    for i, r in enumerate(records):
        ctx = _build_context(r, i + 1, now)
        ws.append([col.format_map(ctx) for col in columns])
    wb.save(path)


def _record_dict(r: ExportRecord, index: int) -> dict:
    return {"index": index, "filename": r.filename, "type": r.type,
            "content": r.content, "count": r.count, "suspect": r.suspect}


def export_json(records: list[ExportRecord], path: str,
                group_by: str = "none") -> None:
    """JSON：结构化数组。none → 记录平铺数组；image/global → 分组嵌套
    [{"group": 键或 null, "items": [记录...]}]。"""
    index = 0
    if group_by == "none":
        payload = [_record_dict(r, (index := index + 1)) for r in records]
    else:
        payload = []
        for key, group in group_records(records, group_by):
            items = [_record_dict(r, (index := index + 1)) for r in group]
            payload.append({"group": key, "items": items})
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
        f.write("\n")


def export(records: list[ExportRecord], path: str, fmt: str,
           template: str = DEFAULT_TEMPLATE, header: str | None = None,
           joiner: str = DEFAULT_JOINER, outer: str = DEFAULT_OUTER,
           group_by: str = "none") -> None:
    fmt = fmt.lower()
    if fmt == "txt":
        export_txt(records, path, template, joiner, outer, group_by)
    elif fmt == "csv":
        export_csv(records, path, template, header)
    elif fmt == "xlsx":
        export_xlsx(records, path, template, header)
    elif fmt == "json":
        export_json(records, path, group_by)
    else:
        raise ValueError(f"不支持的导出格式: {fmt}")
